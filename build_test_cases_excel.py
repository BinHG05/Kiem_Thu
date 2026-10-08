from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).with_name("test-cases-selenium-pom.xlsx")

rows = [
    ["TC01", "Login page", "Login form availability", "Mở trang Login", "Website truy cập được", "Không có", "Mở URL; chờ form", "Trang Login và form hiển thị", "Critical", "Pass", "test_login_page_opens_with_a_visible_form"],
    ["TC02", "Login page", "Initial form state", "Kiểm tra giá trị ban đầu", "Trang Login đã mở", "Không có", "Đọc username và password", "Hai trường đều rỗng", "Normal", "Pass", "test_username_and_password_are_empty_on_first_load"],
    ["TC03", "Login page", "Username input", "Nhập username", "Trang Login đã mở", "student-test", "Nhập username; đọc lại giá trị", "Username nhận đúng dữ liệu", "Normal", "Pass", "test_username_field_accepts_text"],
    ["TC04", "Login page", "Password input", "Nhập password", "Trang Login đã mở", "sample-password", "Nhập password; đọc lại giá trị", "Password nhận đúng dữ liệu", "Normal", "Pass", "test_password_field_accepts_text"],
    ["TC05", "Login page", "Password input", "Kiểm tra che password", "Trang Login đã mở", "Không có", "Đọc thuộc tính type", "type là password", "Critical", "Pass", "test_password_field_is_masked"],
    ["TC06", "Login page", "Login form controls", "Kiểm tra nhãn nút Login", "Trang Login đã mở", "Không có", "Đọc value nút Login", "Nhãn là Đăng nhập", "Minor", "Pass", "test_login_button_has_expected_label"],
    ["TC07", "Login page", "Remember me", "Kiểm tra trạng thái mặc định", "Trang Login đã mở", "Không có", "Đọc trạng thái checkbox", "Checkbox chưa được chọn", "Normal", "Pass", "test_remember_me_is_unselected_on_first_load"],
    ["TC08", "Login page", "Remember me", "Chọn checkbox remember me", "Trang Login đã mở", "Không có", "Click label; đọc trạng thái", "Checkbox được chọn", "Normal", "Pass", "test_remember_me_can_be_selected"],
    ["TC09", "Login page", "Password recovery", "Kiểm tra link quên mật khẩu", "Trang Login đã mở", "Không có", "Đọc href link", "Link trỏ đến /Login/GetPass", "Normal", "Pass", "test_forgot_password_link_targets_password_recovery"],
    ["TC10", "Login page", "External authentication", "Kiểm tra link Google Login", "Trang Login đã mở", "Không có", "Đọc href link Google", "Link chứa accounts.google.com", "Normal", "Pass", "test_google_login_link_targets_google_accounts"],
    ["TC11", "Authentication", "Valid credentials", "Đăng nhập credential hợp lệ", "Có tài khoản UTC trong biến môi trường", "UTC_USER / UTC_PASS", "Mở; nhập username; nhập password; submit", "Rời khỏi /Login", "Blocker", "Fail / Environment", "test_login_success"],
]

wb = Workbook()
summary = wb.active
summary.title = "Test Summary"
cases = wb.create_sheet("Test Cases")
guide = wb.create_sheet("Execution Guide")

navy = "17365D"
blue = "D9EAF7"
green = "E2F0D9"
yellow = "FFF2CC"
red = "FCE4D6"
white = "FFFFFF"
thin = Side(style="thin", color="D9E2F3")
for ws in (summary, cases, guide):
    ws.sheet_view.showGridLines = False

summary.merge_cells("A1:H1")
summary["A1"] = "BUỔI 8 — SELENIUM & POM LOGIN TEST REPORT"
summary["A1"].font = Font(size=16, bold=True, color=white)
summary["A1"].fill = PatternFill("solid", fgColor=navy)
summary["A1"].alignment = Alignment(horizontal="center")
summary.row_dimensions[1].height = 30

info = [
    ("Website", "https://vanphongdientu.utc.edu.vn/Login"),
    ("Framework", "Python + Selenium WebDriver + pytest + POM"),
    ("Browser", "Google Chrome headless"),
    ("Wait strategy", "Explicit Wait — WebDriverWait / Expected Conditions"),
    ("Report", "Allure results + screenshot/page source on failure"),
]
for row, (label, value) in enumerate(info, start=3):
    summary[f"A{row}"] = label
    summary[f"B{row}"] = value
    summary[f"A{row}"].font = Font(bold=True, color=navy)
    summary[f"A{row}"].fill = PatternFill("solid", fgColor=blue)
    summary[f"B{row}"].alignment = Alignment(wrap_text=True)

summary["A10"] = "Chỉ số"
summary["B10"] = "Giá trị"
for cell in summary[10][:2]:
    cell.font = Font(bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
metrics = [
    ("Tổng số test case", "=COUNTA('Test Cases'!A2:A12)"),
    ("Pass theo kế hoạch", '=COUNTIF(\'Test Cases\'!J2:J12,"Pass")'),
    ("Fail / cần credential", '=COUNTIF(\'Test Cases\'!J2:J12,"Fail / Environment")'),
    ("Critical / Blocker", '=COUNTIF(\'Test Cases\'!I2:I12,"Critical")+COUNTIF(\'Test Cases\'!I2:I12,"Blocker")'),
]
for row, (label, formula) in enumerate(metrics, start=11):
    summary[f"A{row}"] = label
    summary[f"B{row}"] = formula
    summary[f"A{row}"].font = Font(bold=True)
    summary[f"B{row}"].font = Font(bold=True, color=navy)
summary["A17"] = "Ghi chú"
summary["B17"] = "TC11 phụ thuộc UTC_USER và UTC_PASS; trạng thái Fail / Environment phản ánh credential hoặc môi trường hiện tại."
summary["B17"].alignment = Alignment(wrap_text=True)
summary.row_dimensions[17].height = 40
summary.column_dimensions["A"].width = 25
summary.column_dimensions["B"].width = 82

headers = ["TC ID", "Module", "Feature", "Test objective", "Preconditions", "Test data", "Test steps", "Expected result", "Severity", "Planned status", "Automated test"]
cases.append(headers)
for row in rows:
    cases.append(row)
cases.freeze_panes = "A2"
table = Table(displayName="LoginTestCases", ref="A1:K12")
table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
cases.add_table(table)
for cell in cases[1]:
    cell.font = Font(bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
cases.row_dimensions[1].height = 30
for row in cases.iter_rows(min_row=2, max_row=12, min_col=1, max_col=11):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    row[0].font = Font(bold=True, color=navy)
    cases.row_dimensions[row[0].row].height = 42
severity = DataValidation(type="list", formula1='"Blocker,Critical,Normal,Minor"')
status = DataValidation(type="list", formula1='"Pass,Fail / Environment,Not Run,In Progress"')
cases.add_data_validation(severity)
cases.add_data_validation(status)
severity.add("I2:I100")
status.add("J2:J100")
cases.conditional_formatting.add("I2:I100", FormulaRule(formula=['I2="Blocker"'], fill=PatternFill("solid", fgColor=red)))
cases.conditional_formatting.add("I2:I100", FormulaRule(formula=['I2="Critical"'], fill=PatternFill("solid", fgColor=yellow)))
cases.conditional_formatting.add("J2:J100", FormulaRule(formula=['J2="Pass"'], fill=PatternFill("solid", fgColor=green)))
cases.conditional_formatting.add("J2:J100", FormulaRule(formula=['J2="Fail / Environment"'], fill=PatternFill("solid", fgColor=red)))
for col, width in {"A": 10, "B": 16, "C": 24, "D": 30, "E": 31, "F": 24, "G": 38, "H": 42, "I": 13, "J": 20, "K": 48}.items():
    cases.column_dimensions[col].width = width

guide.merge_cells("A1:D1")
guide["A1"] = "HƯỚNG DẪN THỰC THI"
guide["A1"].font = Font(size=15, bold=True, color=white)
guide["A1"].fill = PatternFill("solid", fgColor=navy)
guide["A1"].alignment = Alignment(horizontal="center")
guide.append([])
guide.append(["Mục", "Lệnh / cấu hình", "Mục đích", "Ghi chú"])
guide_rows = [
    ["Cài dependency", "python -m pip install -r requirements.txt", "Cài Selenium, pytest và allure-pytest", "Chạy tại thư mục project"],
    ["Credential", '$env:UTC_USER="your_username"; $env:UTC_PASS="your_password"', "Cấu hình cho TC11", "Không commit credential"],
    ["Chạy test", "pytest -v", "Chạy toàn bộ test", "Chrome chạy headless"],
    ["Allure result", "pytest --alluredir=allure-results --clean-alluredir", "Tạo dữ liệu báo cáo", "Kết quả dùng cho Allure"],
    ["Allure UI", "allure serve allure-results", "Mở báo cáo tương tác", "Cần cài Allure Commandline"],
    ["CI smoke", 'pytest -m "not integration"', "Chạy TC01–TC10", "Push / Pull Request"],
    ["CI authenticated", "pytest -m integration", "Chạy TC11", "Cần GitHub Secrets"],
]
for row in guide_rows:
    guide.append(row)
for cell in guide[3]:
    cell.font = Font(bold=True, color=white)
    cell.fill = PatternFill("solid", fgColor=navy)
    cell.alignment = Alignment(horizontal="center", wrap_text=True)
for row in guide.iter_rows(min_row=4, max_row=10, min_col=1, max_col=4):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(bottom=thin)
    guide.row_dimensions[row[0].row].height = 34
for col, width in {"A": 20, "B": 60, "C": 34, "D": 35}.items():
    guide.column_dimensions[col].width = width
guide.freeze_panes = "A4"

wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.save(OUTPUT)
print(str(OUTPUT).encode("ascii", "backslashreplace").decode("ascii"))
